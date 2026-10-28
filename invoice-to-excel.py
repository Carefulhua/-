#!/usr/bin/env python3
"""读取目录下所有发票PDF，汇总金额并生成Excel统计表。

用法:
    python invoice-to-excel.py [目录路径]

如果不指定目录，默认使用当前目录。
生成文件: 发票汇总_YYYYMMDD_HHMMSS.xlsx
"""
import os
import sys
from datetime import datetime

try:
    from pypdf import PdfReader
except ImportError:
    print("错误: 需要安装 pypdf 库。")
    print("请运行: pip install pypdf")
    sys.exit(1)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("错误: 需要安装 openpyxl 库。")
    print("请运行: pip install openpyxl")
    sys.exit(1)


# ---- 复用 invoice_parser 的提取逻辑 ----
from invoice_parser import extract_invoice_info, setup_console


# ---- Excel 生成 ----

THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)

HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
HEADER_FONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
TOTAL_FILL = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
TOTAL_FONT = Font(name='微软雅黑', size=11, bold=True)
BODY_FONT = Font(name='微软雅黑', size=10)
CENTER = Alignment(horizontal='center', vertical='center')
RIGHT = Alignment(horizontal='right', vertical='center')


def process_directory(dir_path):
    """读取目录中所有PDF发票，生成Excel汇总表"""
    if not os.path.isdir(dir_path):
        print(f"错误: 目录不存在: {dir_path}")
        sys.exit(1)

    files = sorted([
        os.path.join(dir_path, f) for f in os.listdir(dir_path)
        if f.lower().endswith('.pdf') and not f.startswith('~')
    ])

    if not files:
        print("目录中没有找到PDF文件。")
        return

    print(f"找到 {len(files)} 个PDF文件，正在提取信息...")
    print('=' * 50)

    records = []
    errors = 0

    for file_path in files:
        fname = os.path.basename(file_path)
        try:
            reader = PdfReader(file_path)
            full_text = ''
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    full_text += text + ' '

            if not full_text.strip():
                print(f"  ⚠️  {fname}: 无法提取文本，跳过")
                errors += 1
                continue

            info = extract_invoice_info(full_text)
            info['item_name'] = info['item_name'] or '未知项目'
            info['total_amount'] = info['total_amount'] or '0'
            records.append(info)
            print(f"  ✓ {info['invoice_number']} | {info['item_name']} | {info['total_amount']}")

        except Exception as e:
            print(f"  ✗ {fname}: {e}")
            errors += 1

    if not records:
        print("没有成功提取任何发票信息。")
        return

    # 生成Excel
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_path = os.path.join(dir_path, f'发票汇总_{timestamp}.xlsx')

    wb = Workbook()
    ws = wb.active
    ws.title = '发票汇总'

    # 列宽
    col_widths = [6, 28, 30, 14]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 标题行
    headers = ['序号', '发票号码', '项目名称', '价税合计']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    # 数据行
    total_sum = 0.0
    for idx, rec in enumerate(records, 1):
        row = idx + 1
        ws.cell(row=row, column=1, value=idx).font = BODY_FONT
        ws.cell(row=row, column=1).alignment = CENTER
        ws.cell(row=row, column=1).border = THIN_BORDER

        ws.cell(row=row, column=2, value=rec['invoice_number']).font = BODY_FONT
        ws.cell(row=row, column=2).border = THIN_BORDER

        ws.cell(row=row, column=3, value=rec['item_name']).font = BODY_FONT
        ws.cell(row=row, column=3).border = THIN_BORDER

        amt_cell = ws.cell(row=row, column=4, value=float(rec['total_amount']))
        amt_cell.font = BODY_FONT
        amt_cell.alignment = RIGHT
        amt_cell.border = THIN_BORDER
        amt_cell.number_format = '#,##0.00'
        total_sum += float(rec['total_amount'])

    # 合计行
    total_row = len(records) + 2
    ws.cell(row=total_row, column=1, value='合计').font = TOTAL_FONT
    ws.cell(row=total_row, column=1).fill = TOTAL_FILL
    ws.cell(row=total_row, column=1).alignment = CENTER
    ws.cell(row=total_row, column=1).border = THIN_BORDER

    for col in [2, 3]:
        c = ws.cell(row=total_row, column=col)
        c.fill = TOTAL_FILL
        c.border = THIN_BORDER

    sum_cell = ws.cell(row=total_row, column=4, value=total_sum)
    sum_cell.font = TOTAL_FONT
    sum_cell.fill = TOTAL_FILL
    sum_cell.alignment = RIGHT
    sum_cell.border = THIN_BORDER
    sum_cell.number_format = '#,##0.00'

    # 统计信息行
    info_row = total_row + 2
    ws.cell(row=info_row, column=1, value=f'共 {len(records)} 张发票').font = BODY_FONT
    if errors:
        ws.cell(row=info_row + 1, column=1, value=f'失败 {errors} 张').font = BODY_FONT

    wb.save(output_path)
    print()
    print('=' * 50)
    print(f"✅ 已生成汇总表: {output_path}")
    print(f"   共 {len(records)} 张发票，合计金额: {total_sum:.2f} 元")
    if errors:
        print(f"   ⚠️  {errors} 张发票读取失败")
    print('=' * 50)


if __name__ == '__main__':
    setup_console()

    if len(sys.argv) >= 2:
        target_dir = sys.argv[1]
    else:
        target_dir = os.getcwd()

    print('=' * 50)
    print('   发票PDF → Excel 汇总工具')
    print('=' * 50)
    print()

    process_directory(target_dir)
